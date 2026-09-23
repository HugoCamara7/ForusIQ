"""Archivo formato Neogística: nombres, orden y formato tomados del reporte original."""

import io
import json

import openpyxl
import pandas as pd
import pytest

from forusight.data.github_store import GitHubStore, config_github
from forusight.data.repository import SinAlmacenamiento, guardar_aprobacion_github
from forusight.data.synthetic import generar
from forusight.engine.pipeline import ejecutar
from forusight.export.neogistica import a_excel_neogistica, construir_tabla, nombre_archivo
from tests.builders import params

#: Encabezado real de "534 - Sugerido de Distribución (Extendido)" (sin las 12 semanas).
NEOGISTICA = [
    "Código SKU",
    "Modelo",
    "Color",
    "Código Modelo",
    "Código Color",
    "Talla",
    "Descripción SKU",
    "Clase",
    "Marca",
    "Género",
    "Prenda",
    "Temporada comercial",
    "Mantener Inventario",
    "Reposición Bloqueada",
    "Código Centro",
    "Nombre Centro",
    "Centro Comercial",
    "Zona CC",
    "Código Grupo Planificación",
    "Grupo Requerimiento",
    "Costo [$/un]",
    "Costo Ponderado [$/un]",
    "Precio Venta [$/un]",
    "Clase Demanda Frecuencia",
    "Clase Demanda Variabilidad",
    "Clase Obsolescencia",
    "SEMANAS",
    "Demanda Periodo Actual",
    "Pronóstico Demanda [un/semana]",
    "Leadtime [días]",
    "Período Revisión [días]",
    "Nivel de Disponibilidad Planificado [%]",
    "Punto Reorden [un]",
    "Nivel Máximo [un]",
    "Stock Mínimo Total",
    "Código Centro Origen",
    "Stock en CD",
    "Stock Físico [un]",
    "Stock Trán. Int. [un]",
    "Stock Trán. Prov. [un]",
    "Stock Comprometido [un]",
    "Backorder [un]",
    "Posición Stock [un]",
    "Sobrestock Multiempresa [un]",
    "Cantidad Pedida Final [un]",
    "Pendiente Reposición",
    "Motivo Pendiente Reposición",
    "Unidad Empaque Distribución",
    "Unidad Venta",
    "Alcance Posición Stock Actual [semanas]",
    "Alcance Posición Stock Final [semanas]",
    "Monto Pedido Final [$]",
]
MOTIVOS = {"Sin Reposición Pendiente", "Stock CD", "Almacenamiento"}


@pytest.fixture(scope="module")
def corrida():
    inp, corte = generar()
    p = params()
    return inp, corte, p, ejecutar(inp, p, corte, run_id="R")


def _tabla(corrida, cantidad=None):
    inp, corte, p, r = corrida
    return construir_tabla(
        r.detalle, inp.ventas, inp.dim_producto, inp.dim_tienda, p, corte, cantidad=cantidad
    )


def test_columnas_son_del_reporte_y_en_su_orden(corrida):
    t = _tabla(corrida)
    semanas = [c for c in t.columns if c[:2] == "20"]
    assert len(semanas) == 12 and semanas == sorted(semanas)
    nombres = ["SEMANAS" if c in semanas else c for c in t.columns]
    nombres = [n for i, n in enumerate(nombres) if n != "SEMANAS" or nombres.index("SEMANAS") == i]
    assert set(nombres) <= set(NEOGISTICA), set(nombres) - set(NEOGISTICA)
    posiciones = [NEOGISTICA.index(n) for n in nombres]
    assert posiciones == sorted(posiciones)
    # no se inventa lo que no hay: costos, clases de demanda, backorder…
    for omitida in (
        "Costo [$/un]",
        "Clase Demanda Frecuencia",
        "Backorder [un]",
        "Monto Pedido Final [$]",
        "Punto Reorden [un]",
    ):
        assert omitida not in t.columns


def test_valores_operativos(corrida):
    t = _tabla(corrida)
    assert set(t["Motivo Pendiente Reposición"]) <= MOTIVOS
    assert (
        t.loc[t["Pendiente Reposición"] > 0, "Motivo Pendiente Reposición"]
        != "Sin Reposición Pendiente"
    ).all()
    assert (t["Código Centro Origen"] == "320").all()
    pos = t["Stock Físico [un]"] + t.get("Stock Trán. Int. [un]", 0)
    assert (t["Posición Stock [un]"] == pos).all()
    assert set(t["Grupo Requerimiento"]) <= {"Revision de stock", "Carga Pedidos"}
    # sólo filas con actividad (resumido)
    act = (
        (t["Cantidad Pedida Final [un]"] > 0)
        | (t["Pendiente Reposición"] > 0)
        | (t["Stock Físico [un]"] > 0)
        | (t[[c for c in t if c[:2] == "20"]].sum(axis=1) > 0)
        | (t.get("Stock Trán. Int. [un]", 0) > 0)
    )
    assert act.all()


def test_cantidad_aprobada_reemplaza_la_propuesta(corrida):
    _, _, _, r = corrida
    cero = pd.Series(0, index=r.detalle.index)
    t = _tabla(corrida, cantidad=cero)
    assert t["Cantidad Pedida Final [un]"].sum() == 0


def test_excel_con_cabecera_y_estilo_neogistica(corrida):
    t = _tabla(corrida)
    x = a_excel_neogistica(t, pd.Timestamp("2026-09-23"))
    wb = openpyxl.load_workbook(io.BytesIO(x))
    ws = wb["Hoja1"]
    assert ws["A3"].value == "Empresa:" and ws["B3"].value == "Forus Peru"
    assert ws["A4"].value == "Reporte:" and ws["A5"].value == "Fecha:"
    assert ws["B5"].value == "23/09/2026"
    assert ws["A7"].value == "Código SKU" and ws.auto_filter.ref.startswith("A7")
    assert ws["A7"].fill.fgColor.rgb.endswith("084B8A") and ws["A7"].font.b
    j = list(t.columns).index("Cantidad Pedida Final [un]") + 1
    celda = ws.cell(8, j)
    assert celda.fill.fgColor.rgb.endswith("C2C567") and celda.font.color.rgb.endswith("FF0000")
    assert ws.max_row == 7 + len(t)
    assert "Resumen" in wb.sheetnames
    assert nombre_archivo(pd.Timestamp("2026-09-23")).startswith("20260923_Reporte_Distribucion")


def test_config_github_reutiliza_ticketing_del_catalogo():
    assert config_github({}) is None
    tk = {
        "ticketing": {
            "backend": "github",
            "repository": "HugoCamara7/datos",
            "token": "t",
            "branch": "catalog-tickets",
        }
    }
    c = config_github(tk)
    assert c == {
        "repository": "HugoCamara7/datos",
        "token": "t",
        "branch": "catalog-tickets",
        "prefix": "forusight",
    }
    assert (
        config_github({"ticketing": {"repository": "a/b", "token": "GITHUB_TOKEN_CON_PERMISO"}})
        is None
    )


def test_github_store_hace_put_con_sha(monkeypatch):
    llamadas = []

    class Resp:
        def __init__(self, data):
            self.data = data

        def read(self):
            return json.dumps(self.data).encode()

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    def opener(req, timeout=30):
        llamadas.append((req.get_method(), req.full_url, req.data))
        return Resp({"sha": "abc"} if req.get_method() == "GET" else {"content": {"sha": "n"}})

    s = GitHubStore("o/r", "t", "rama", "forusight", opener=opener)
    assert s.guardar("aprobaciones/x.csv", b"a,b", "msg") == "forusight/aprobaciones/x.csv"
    assert llamadas[0][0] == "GET" and "ref=rama" in llamadas[0][1]
    body = json.loads(llamadas[1][2])
    assert llamadas[1][0] == "PUT" and body["sha"] == "abc" and body["branch"] == "rama"


def test_sin_dataset_ni_github_la_aprobacion_pide_descarga():
    with pytest.raises(SinAlmacenamiento, match="Descárgala"):
        guardar_aprobacion_github("R", pd.DataFrame({"a": [1]}), "u", secrets={})
